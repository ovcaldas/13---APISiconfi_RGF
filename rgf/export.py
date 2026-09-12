from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import DATA_DIR


def highlights(indicators: pd.DataFrame) -> dict[str, pd.DataFrame]:
    return {
        "Melhores": indicators.nsmallest(5, "ranking_historico"),
        "Piores": indicators.nlargest(5, "ranking_historico"),
        "Alertas": indicators.sort_values(["anos_acima_maximo", "anos_acima_prudencial", "anos_acima_alerta"], ascending=False),
    }


def executive_summary(base: pd.DataFrame, indicators: pd.DataFrame) -> str:
    latest_year = int(base.dropna(subset=["DTP_RCL"])["ano"].max())
    latest = base[(base["ano"] == latest_year) & base["DTP_RCL"].notna()]
    best = ", ".join(indicators.nsmallest(5, "ranking_historico")["UF"])
    worst = ", ".join(indicators.nlargest(5, "ranking_historico")["UF"])
    pressured = latest[latest["situacao_fiscal"].isin(["alerta", "acima do limite prudencial", "acima do limite máximo"])]
    improving = ", ".join(indicators.query("tendencia_recente == 'melhora'").nsmallest(5, "tendencia_recente_pp_ano")["UF"])
    worsening = ", ".join(indicators.query("tendencia_recente == 'deterioração'").nlargest(5, "tendencia_recente_pp_ano")["UF"])
    region = latest.groupby("regiao")["DTP_RCL"].mean().sort_values()
    region_text = "; ".join(f"{idx}: {value:.2f}%" for idx, value in region.items())
    return f"""Panorama da Despesa com Pessoal dos Estados Brasileiros – RGF 2015–2025

Escopo: Poder Executivo estadual, último quadrimestre disponível de cada exercício. No último ano disponível ({latest_year}), há {len(latest)} UFs com indicador válido. A leitura usa os limites declarados no próprio Anexo 01 do RGF.

Os cinco melhores resultados no score multidimensional são {best}; os cinco piores são {worst}. Em {latest_year}, {len(pressured)} UFs estavam em alerta ou acima de limite prudencial/máximo. As médias regionais do DTP/RCL foram: {region_text}.

Tendência recente: melhora mais pronunciada em {improving or 'nenhuma UF classificada'}; deterioração mais pronunciada em {worsening or 'nenhuma UF classificada'}. Os principais riscos para os próximos exercícios concentram-se nas UFs com margem pequena ou negativa, recorrência de ultrapassagens, tendência crescente e alta volatilidade. O score é comparativo e não substitui a avaliação legal de cada demonstrativo homologado.
"""


def export_excel(base: pd.DataFrame, indicators: pd.DataFrame, ranking: pd.DataFrame) -> Path:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    target = DATA_DIR / "rgf_estados_2015_2025.xlsx"
    latest_year = int(base.dropna(subset=["DTP_RCL"])["ano"].max())
    latest = ranking[ranking["ano"] == latest_year]
    parts = highlights(indicators)
    resumo = indicators[["UF", "nome_estado", "regiao", "score_fiscal", "ranking_historico", "grupo", "tendencia_recente", "media_2015_2025", "media_ultimos_3_anos", "anos_acima_alerta", "anos_acima_prudencial", "anos_acima_maximo"]]
    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        base.to_excel(writer, sheet_name="Base_RGF", index=False)
        indicators.to_excel(writer, sheet_name="Indicadores", index=False)
        latest.to_excel(writer, sheet_name="Ranking_2025", index=False)
        indicators.sort_values("ranking_historico").to_excel(writer, sheet_name="Ranking_Historico", index=False)
        parts["Melhores"].to_excel(writer, sheet_name="Melhores", index=False)
        parts["Piores"].to_excel(writer, sheet_name="Piores", index=False)
        parts["Alertas"].to_excel(writer, sheet_name="Alertas", index=False)
        resumo.to_excel(writer, sheet_name="Resumo_Estados", index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="19324D"); cell.font = Font(color="FFFFFF", bold=True); cell.alignment = Alignment(wrap_text=True)
            for col_idx, cells in enumerate(ws.columns, 1):
                width = min(max(len(str(c.value or "")) for c in cells) + 2, 38)
                ws.column_dimensions[get_column_letter(col_idx)].width = width
            headers = {cell.value: cell.column for cell in ws[1]}
            for score_header in ("score_fiscal", "DTP_RCL"):
                if score_header in headers and ws.max_row > 1:
                    letter = get_column_letter(headers[score_header])
                    ws.conditional_formatting.add(f"{letter}2:{letter}{ws.max_row}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    return target

